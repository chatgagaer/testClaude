# -*- coding: utf-8 -*-
"""
读取 Excel 中的 text-cn 字段，逐行调用 HTTP 接口，
将响应 JSON 中的 result 写入同一行的 result-cn-q 列。

依赖: pip install openpyxl requests
用法: python call_api_fill_excel.py
"""

# ==================== 配置区（请自行填写） ====================
API_URL = ""                       # 接口地址，例如 "http://127.0.0.1:8080/api/check"
HEADERS = {"Content-Type": "application/json"}  # 如需鉴权可在此添加 token 等请求头
TIMEOUT = 30                       # 请求超时时间（秒）

INPUT_FILE = "test300V0.6.xlsx"    # 输入 Excel 文件路径
OUTPUT_FILE = INPUT_FILE           # 输出文件路径，默认覆盖原文件；也可改成新文件名
SHEET_NAME = "数据测试集"           # 工作表名称，填 None 则使用第一个工作表

SOURCE_COLUMN = "text-cn"          # 读取的列名
TARGET_COLUMN = "result-cn-q"      # 写入的列名
CONTENT_TYPE = "question"          # 请求体中写死的 contentType
SKIP_FILLED = True                 # True: 已有结果的行跳过（便于中断后续跑）
SAVE_EVERY = 20                    # 每处理多少行保存一次
# ============================================================

import json
import sys

import openpyxl
import requests


def find_column(ws, name):
    """在表头（第一行）中查找列名，返回列号（从 1 开始）。"""
    for cell in ws[1]:
        if cell.value is not None and str(cell.value).strip() == name:
            return cell.column
    raise ValueError(f"表头中未找到列: {name}")


def call_api(text):
    """调用接口，成功返回 result 内容，失败返回错误描述字符串。"""
    payload = {"content": text, "contentType": CONTENT_TYPE}
    try:
        resp = requests.post(API_URL, json=payload, headers=HEADERS, timeout=TIMEOUT)
    except requests.RequestException as e:
        return f"ERROR: 请求异常 {e}"

    if resp.status_code != 200:
        return f"ERROR: HTTP {resp.status_code}"

    try:
        data = resp.json()
    except ValueError:
        return f"ERROR: 响应不是合法 JSON: {resp.text[:200]}"

    result = data.get("result") if isinstance(data, dict) else None
    # result 若为对象/数组，则转为 JSON 字符串写入单元格
    if isinstance(result, (dict, list)):
        return json.dumps(result, ensure_ascii=False)
    return result


def main():
    if not API_URL:
        sys.exit("请先在文件上方填写 API_URL")

    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.worksheets[0]

    src_col = find_column(ws, SOURCE_COLUMN)
    dst_col = find_column(ws, TARGET_COLUMN)

    processed = 0
    for row in range(2, ws.max_row + 1):
        text = ws.cell(row=row, column=src_col).value
        if text is None or str(text).strip() == "":
            continue
        target = ws.cell(row=row, column=dst_col)
        if SKIP_FILLED and target.value not in (None, ""):
            continue

        result = call_api(str(text))
        target.value = result
        processed += 1
        print(f"[第{row}行] {str(text)[:30]} -> {str(result)[:50]}")

        if processed % SAVE_EVERY == 0:
            wb.save(OUTPUT_FILE)

    wb.save(OUTPUT_FILE)
    print(f"完成，共处理 {processed} 行，已保存到 {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
